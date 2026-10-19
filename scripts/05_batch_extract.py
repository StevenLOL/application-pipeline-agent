"""scripts/05_batch_extract.py — F3 batch extraction + agreement table + ledger cross-validation (D10).

Responsibilities (docs/10 D10 acceptance: "extract from 10+ historical JDs,
produce the agreement table"):
1. Scan singapore/jobs/投递包/*/JD.md (read-only), take the "## JD 原文" section
   and the record ID from the header;
2. For each JD: deadline runs 5-run (measured: the only reliably extractable
   target); requirements/assessments run 1-run each (known 1.5B limit — only
   the schema status is recorded — reports/f3/run_001.md);
3. **Ledger cross-validation**: extracted deadline vs the ledger's "截止日期"
   (read-only openpyxl) — three states: match / mismatch / ledger value
   missing; this is the substance of "F3 aligns with the ledger schema": the
   extracted field semantics line up with the ledger columns
   (record ID / deadline);
4. needs_human rows land in `data/extraction_review.json` (low-confidence
   human review queue, docs/08 §6);
5. Print the agreement table + write markdown to `reports/f3/run_002.md`.

Usage (repo root, conda env apa, llama-server on 8080):
  python scripts/05_batch_extract.py [--limit N] [--out reports/f3/run_002.md]
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from apa.extract import JDExtractor, configure_from_env, normalize_date  # noqa: E402

_SG = Path(__import__("os").environ.get("APA_SINGAPORE_ROOT",
                                        r"<private-workspace>"))
PACKAGES_DIR = _SG / "jobs" / "投递包"
LEDGER_XLSX = _SG / "docs" / "JD总台账.xlsx"
REVIEW_QUEUE = ROOT / "data" / "extraction_review.json"

_JD_SECTION_RE = re.compile(r"## JD 原文\s*\n(.*)", re.S)
_RECORD_ID_RE = re.compile(r"\*\*记录 ID\*\*\s*\|\s*`([^`]+)`")
_DEADLINE_HEADER_RE = re.compile(r"\*\*🔴 截止日期\*\*\s*\|\s*\*\*([^*]+)\*\*")


def collect_jds(limit: int) -> list[dict]:
    """Scan the delivery packages; return [{jd_id, path, jd_text, header_deadline}]."""
    out = []
    for jd_path in sorted(PACKAGES_DIR.glob("*/JD.md")):
        text = jd_path.read_text(encoding="utf-8", errors="replace")
        m = _JD_SECTION_RE.search(text)
        if not m:
            continue
        rid = _RECORD_ID_RE.search(text)
        dl = _DEADLINE_HEADER_RE.search(text)
        out.append({
            "jd_id": rid.group(1) if rid else jd_path.parent.name,
            "path": str(jd_path),
            "jd_text": m.group(1).strip(),
            "header_deadline": dl.group(1).strip() if dl else "",
        })
        if len(out) >= limit:
            break
    return out


def load_ledger_deadlines() -> dict[str, str]:
    """Ledger (read-only): record ID -> raw deadline text."""
    import openpyxl

    if not LEDGER_XLSX.exists():
        return {}
    wb = openpyxl.load_workbook(LEDGER_XLSX, data_only=True)
    ws = wb["总台账"]
    hmap = {ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)}
    out = {}
    for r in range(2, ws.max_row + 1):
        rid = str(ws.cell(row=r, column=hmap["记录ID"]).value or "")
        out[rid] = str(ws.cell(row=r, column=hmap["截止日期"]).value or "")
    return out


def compare_with_ledger(extracted: str, ledger_value: str) -> str:
    """Three states comparing the extracted value vs the ledger value:
    match / mismatch / ledger-missing.

    Ledger deadlines have day granularity or are "page not published"-style
    text; the extracted value is YYYY-MM-DD or NOT_PUBLISHED. The comparison
    happens after normalization (if the ledger value includes a time, take
    the date part).
    """
    ledger = (ledger_value or "").strip()
    if not ledger or ledger in ("-", "None"):
        return "ledger-missing"
    if extracted == "NOT_PUBLISHED":
        # ledger text like "page not published" counts as a match
        return "match" if ("未公布" in ledger or "WORKABLE" in ledger.upper()) else "mismatch"
    led_norm = normalize_date(ledger) or ledger.split(" ")[0]
    return "match" if led_norm == extracted else "mismatch"


def main() -> int:
    ap = argparse.ArgumentParser(description="F3 批量抽取 + 一致率表")
    ap.add_argument("--limit", type=int, default=25)
    ap.add_argument("--out", default=str(ROOT / "reports" / "f3" / "run_002.md"))
    args = ap.parse_args()

    jds = collect_jds(args.limit)
    if len(jds) < 10:
        print(f"error: 仅找到 {len(jds)} 份 JD（验收要求 10+）", file=sys.stderr)
        return 2
    print(f"收集 {len(jds)} 份 JD", flush=True)

    lm = configure_from_env(backend="local", env={
        "APA_LLM_BASE_URL": "http://127.0.0.1:8080/v1",
        "APA_LLM_MODEL": "local", "APA_LLM_API_KEY": "EMPTY",
    })
    extractor = JDExtractor(lm=lm)
    ledger = load_ledger_deadlines()

    rows: list[dict] = []
    for i, jd in enumerate(jds, 1):
        print(f"[{i}/{len(jds)}] {jd['jd_id']}", flush=True)
        result = extractor.extract_with_consistency(jd["jd_text"], jd_ref=jd["jd_id"], n=5)
        dl_cons = result.consistency["deadline"]
        extracted_dl = result.values.get("deadline")
        cmp_state = compare_with_ledger(extracted_dl, ledger.get(jd["jd_id"], "")) if extracted_dl else "extract-failed"
        rows.append({
            "jd_id": jd["jd_id"],
            "header_deadline": jd["header_deadline"],
            "ledger_deadline": ledger.get(jd["jd_id"], ""),
            "extracted_deadline": extracted_dl,
            "deadline_agreement": round(dl_cons.agreement, 2),
            "deadline_valid_runs": f"{dl_cons.valid_runs}/5",
            "req_schema": "ok" if "requirements" not in result.schema_errors else "fail(1.5B 边界)",
            "asmt_schema": "ok" if "assessments" not in result.schema_errors else "fail(1.5B 边界)",
            "asmt_dropped": len(result.dropped_assessments),
            "ledger_compare": cmp_state,
            "needs_human": result.needs_human,
            "schema_errors": dict(result.schema_errors),
        })

    # agreement table
    n_match = sum(1 for r in rows if r["ledger_compare"] == "match")
    n_mismatch = sum(1 for r in rows if r["ledger_compare"] == "mismatch")
    n_missing = sum(1 for r in rows if r["ledger_compare"] == "ledger-missing")
    n_full_agree = sum(1 for r in rows if r["deadline_agreement"] == 1.0)
    n_needs_human = sum(1 for r in rows if r["needs_human"])

    lines = [
        "# F3 run_002：批量抽取一致率表（D10）",
        "",
        f"> 时间：{datetime.now().strftime('%Y-%m-%d %H:%M')} | JD 数：{len(rows)}"
        f" | 模型：Qwen2.5-1.5B local（D9 已知边界：列表目标仅记录 schema 状态）",
        "",
        "## 汇总",
        "",
        f"- deadline 5-run 满一致率：**{n_full_agree}/{len(rows)}**",
        f"- deadline vs 台账：match **{n_match}** / mismatch **{n_mismatch}**"
        f" / ledger-missing **{n_missing}**",
        f"- needs_human（低置信转人工）：**{n_needs_human}/{len(rows)}**"
        f"（队列：`data/extraction_review.json`）",
        "",
        "## 明细",
        "",
        "| 记录ID | 抽取deadline | 台账 | 比对 | 一致率 | schema | needs_human |",
        "|---|---|---|---|---|---|---|",
    ]
    for r in rows:
        lines.append(
            f"| {r['jd_id']} | {r['extracted_deadline'] or '-'} "
            f"| {r['ledger_deadline'][:20] or '-'} | {r['ledger_compare']} "
            f"| {r['deadline_agreement']} ({r['deadline_valid_runs']}) "
            f"| req:{r['req_schema']}/asmt:{r['asmt_schema']}"
            f"{'/drop:' + str(r['asmt_dropped']) if r['asmt_dropped'] else ''} "
            f"| {'Y' if r['needs_human'] else 'N'} |"
        )
    lines += [
        "",
        "## 判读纪律",
        "",
        "- mismatch ≠ 抽取错：台账可能过期/粒度不同（页面 27 Oct vs 台账 22 Oct 的"
        " Contoso 案例）——**逐条人工复核后经 `jd_registry_ops.py` 更新**，抽取只提供线索；",
        "- ledger-missing 是 docs/13 🔴「截止列 7.1%」的直接量化——抽取层正是"
        " 补齐该列的候选数据源（经沙箱验证后人工确认）；",
        "- needs_human 的完整上下文（schema 错误）入 `data/extraction_review.json`。",
    ]

    out_path = Path(args.out)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text("\n".join(lines), encoding="utf-8")
    print(f"\n报告：{out_path}")
    print(f"汇总：满一致率 {n_full_agree}/{len(rows)}；match {n_match}/"
          f"mismatch {n_mismatch}/missing {n_missing}；needs_human {n_needs_human}")

    # low-confidence human review queue
    REVIEW_QUEUE.parent.mkdir(parents=True, exist_ok=True)
    review = [
        {"jd_id": r["jd_id"], "schema_errors": r["schema_errors"],
         "ledger_compare": r["ledger_compare"]}
        for r in rows if r["needs_human"]
    ]
    REVIEW_QUEUE.write_text(
        json.dumps({"generated": datetime.now().isoformat(timespec="seconds"),
                    "items": review}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(f"人工队列：{REVIEW_QUEUE}（{len(review)} 项）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
