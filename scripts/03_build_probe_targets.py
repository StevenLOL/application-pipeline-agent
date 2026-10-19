"""D4: extract active job URLs from the JD master ledger -> data/probe_targets.json
(target source for the F10 JobUrlPoller).

Data flow: singapore JD总台账.xlsx (single source of truth, docs/116 iron rule:
never modify) -> this script extracts read-only -> data/probe_targets.json
(gitignored: delivery strategy is personal data).

Filtering criteria (conservative: prefer fewer targets over noise):
  - verdict in {可投, 可投(待用户裁决)}: only jobs with real delivery intent
    are worth a 15-min probe; 未评/备选/待核 are not probed (no delivery
    decision yet); 不投/已关闭/已投递/已投 are not probed (nothing in flight)
  - deadline empty or >= today (SGT): probing a past-deadline job's URL to a
    200 is just noise (measured 2026-10-18: a past-deadline Initech job still
    returned HEAD 200 — MCF does not take down closed pages)
  - real link starts with http(s); deduplicate by URL

Ledger hygiene signal: rows already past deadline but still verdicted 可投 are
counted in the summary (for the user to clean the ledger; the ledger is never
modified automatically — iron rule: the only channel for JD CRUD is
jd_registry_ops.py).

Usage:
    conda run -n apa python scripts/03_build_probe_targets.py [--ledger PATH] [--out PATH]
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parent.parent
import os as _os

_SG = _os.environ.get("APA_SINGAPORE_ROOT", r"<private-workspace>")
DEFAULT_LEDGER = Path(_SG) / "docs" / "JD总台账.xlsx"
DEFAULT_OUT = REPO_ROOT / "data" / "probe_targets.json"

ACTIVE_VERDICTS = {"可投", "可投(待用户裁决)"}
SHEET = "总台账"

# Valid URL head (ledger cells have contained dirty annotated values like
# "https://www.proseware.example (x0pa API)": take the leading URL and keep
# the annotation as a note instead of dropping it — probes must be
# requestable, information must stay traceable)
_URL_RE = re.compile(r"^(https?://[^\s]+)")

# Channel classification: under the jobs.careers.gov.sg host, split by ATS
# backend path — the status page treats HRPS / Greenhouse / Workable as
# separate components (uptime is measured independently); lumping them
# together would book a single-backend outage against the whole portal
# (fixed 2026-10-18: workable jobs used to be classified as hrp)
HOST_CHANNEL = [
    ("mycareersfuture.gov.sg", "mcf"),
    ("jobs.careers.gov.sg", "hrp"),          # default backend; refined by path below
]


def classify_channel(url: str, source: str) -> str:
    low = url.lower()
    for host, ch in HOST_CHANNEL:
        if host in low:
            if ch == "hrp":
                if "/greenhouse/" in low:
                    return "hrp-greenhouse"
                if "/workable/" in low:
                    return "hrp-workable"
            return ch
    src = (source or "").upper()
    if "GREENHOUSE" in src:
        return "hrp-greenhouse"
    if "WORKABLE" in src:
        return "hrp-workable"
    if "HRP" in src or "GOV" in src:
        return "hrp"
    if "MCF" in src:
        return "mcf"
    return "other"


def parse_deadline(value) -> date | None:
    """Lenient deadline parsing: datetime / several date string formats;
    returns None on failure (treated as unknown)."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    s = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d %b %Y", "%d %B %Y", "%Y/%m/%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def build(ledger_path: Path) -> dict:
    wb = openpyxl.load_workbook(ledger_path, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        raise SystemExit(f"台账缺少工作表 {SHEET!r}（现有: {wb.sheetnames()}）")
    ws = wb[SHEET]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h else "" for h in next(rows)]
    idx = {name: i for i, name in enumerate(header)}

    need = ["记录ID", "真实链接", "裁决", "截止日期", "岗位标题", "雇主/机构", "来源"]
    missing = [c for c in need if c not in idx]
    if missing:
        raise SystemExit(f"台账缺少必需列: {missing}（现有: {header}）")

    today = date.today()
    targets: dict[str, dict] = {}
    stats = {"total": 0, "active_verdict": 0, "past_deadline": 0,
             "no_url": 0, "deadline_unknown": 0, "dupe_url": 0,
             "malformed_url_fixed": 0}

    for row in rows:
        if row is None or all(c in (None, "") for c in row):
            continue
        stats["total"] += 1
        verdict = str(row[idx["裁决"]] or "").strip()
        if verdict not in ACTIVE_VERDICTS:
            continue
        stats["active_verdict"] += 1
        raw_url = str(row[idx["真实链接"]] or "").strip()
        m = _URL_RE.match(raw_url)
        if not m:
            stats["no_url"] += 1
            continue
        url = m.group(1).rstrip(".,;)")     # trailing punctuation is not part of the URL
        url_note = raw_url[len(m.group(1)):].strip()
        if url_note:
            stats["malformed_url_fixed"] += 1
        dl = parse_deadline(row[idx["截止日期"]])
        if dl is None:
            stats["deadline_unknown"] += 1
        elif dl < today:
            stats["past_deadline"] += 1
            continue
        if url in targets:
            stats["dupe_url"] += 1
            continue
        target = {
            "jd_id": str(row[idx["记录ID"]] or ""),
            "url": url,
            "title": str(row[idx["岗位标题"]] or "")[:120],
            "employer": str(row[idx["雇主/机构"]] or "")[:80],
            "channel": classify_channel(url, str(row[idx["来源"]] or "")),
            "deadline": dl.isoformat() if dl else None,
            "verdict": verdict,
        }
        if url_note:
            target["url_note"] = url_note[:120]
        targets[url] = target

    return {"generated_at": datetime.now().isoformat(timespec="seconds"),
            "source": str(ledger_path), "today": today.isoformat(),
            "stats": stats, "targets": sorted(targets.values(),
                                              key=lambda t: (t["channel"], t["jd_id"]))}


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="JD 台账 → F10 探测目标 JSON")
    ap.add_argument("--ledger", type=Path, default=DEFAULT_LEDGER)
    ap.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = ap.parse_args(argv)

    if not args.ledger.exists():
        print(f"error: 台账不存在: {args.ledger}", file=sys.stderr)
        return 2
    data = build(args.ledger)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")

    s = data["stats"]
    print(f"台账 {s['total']} 行 → 活跃裁决 {s['active_verdict']} 行 → "
          f"探测目标 {len(data['targets'])} 个 URL")
    print(f"  剔除: 已截止 {s['past_deadline']} / 无链接 {s['no_url']} / "
          f"重复 URL {s['dupe_url']}；截止未知(保留) {s['deadline_unknown']}；"
          f"脏链接已剥离注释 {s['malformed_url_fixed']}")
    by_ch: dict[str, int] = {}
    for t in data["targets"]:
        by_ch[t["channel"]] = by_ch.get(t["channel"], 0) + 1
    print(f"  渠道分布: {by_ch}")
    print(f"写出: {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
